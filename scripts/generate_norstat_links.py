#!/usr/bin/env python3
"""Generate unique Norstat participant links with pid query parameters.

Example:
  python3 scripts/generate_norstat_links.py \
    --base-url https://your-ucloud-root.example.com/ \
    --start-id 1001 \
    --count 50000 \
    --output docs/links/norstat_links.xlsx
"""
from __future__ import annotations

import argparse
from pathlib import Path
from urllib.parse import parse_qsl, urlencode, urlsplit, urlunsplit

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


def link_with_pid(base_url: str, pid: int) -> str:
    split = urlsplit(base_url)
    query = dict(parse_qsl(split.query, keep_blank_values=True))
    query["pid"] = str(pid)
    return urlunsplit((split.scheme, split.netloc, split.path or "/", urlencode(query), split.fragment))


def main() -> None:
    parser = argparse.ArgumentParser(description="Generate Norstat unique-link Excel file.")
    parser.add_argument("--base-url", required=True, help="Root study URL, e.g. https://ucloud.example/path or http://localhost:3000/")
    parser.add_argument("--start-id", type=int, default=1001)
    parser.add_argument("--count", type=int, default=50_000)
    parser.add_argument("--output", type=Path, default=Path("docs/ethics/norstat_unique_links.xlsx"))
    args = parser.parse_args()

    if args.count < 1:
        raise SystemExit("--count must be positive")
    args.output.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook(write_only=False)
    ws = wb.active
    ws.title = "Norstat links"
    ws.append(["id", "link"])
    for pid in range(args.start_id, args.start_id + args.count):
        ws.append([pid, link_with_pid(args.base_url, pid)])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 96
    wb.save(args.output)
    print(f"Wrote {args.output} with {args.count} links starting at {args.start_id}")


if __name__ == "__main__":
    main()
